import os
import sys
import datetime
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

folder = 'danang'
files = []
for f in os.listdir(folder):
    fp = os.path.join(folder, f)
    if os.path.isfile(fp):
        stat = os.stat(fp)
        ctime = datetime.datetime.fromtimestamp(stat.st_ctime)
        mtime = datetime.datetime.fromtimestamp(stat.st_mtime)
        files.append({
            'name': f,
            'size': stat.st_size,
            'ctime': ctime,
            'mtime': mtime,
            'ctime_str': ctime.strftime('%Y-%m-%d %H:%M:%S'),
            'mtime_str': mtime.strftime('%Y-%m-%d %H:%M:%S')
        })

# Sắp xếp theo thứ tự thời gian tải về (Creation Time)
files_by_ctime = sorted(files, key=lambda x: x['ctime'])

print(f"=== TỔNG CỘNG CÓ {len(files_by_ctime)} FILE TRONG THƯ MỤC {folder}/ ===")
print("Thứ tự sắp xếp theo THỜI GIAN TẢI VỀ (Creation Time / Sớm nhất -> Mới nhất):\n")

for i, item in enumerate(files_by_ctime, 1):
    sz = item['size']
    sz_str = f"{sz/(1024*1024):.2f} MB" if sz >= 1024*1024 else f"{sz/1024:.1f} KB"
    print(f"{i:2d}. [{item['ctime_str']}] {item['name']} ({sz_str})")

print("\n=== ĐỌC FILE EXCEL ĐÀ NẴNG HIỆN CÓ ===")
wb = openpyxl.load_workbook('Danh_Sach_45_Dia_Diem_Du_Lich_Da_Nang.xlsx')
ws = wb.active
print("Tiêu đề sheet:", ws.title)
print("Tổng số hàng:", ws.max_row)
print("Headers hàng 3:", [ws.cell(3, c).value for c in range(1, ws.max_column+1) if ws.cell(3, c).value is not None])
print("\nDanh sách các địa điểm trong Excel Đà Nẵng:")
for r in range(4, ws.max_row+1):
    stt = ws.cell(r, 1).value
    name = ws.cell(r, 2).value
    cat = ws.cell(r, 3).value
    district = ws.cell(r, 4).value
    print(f"{stt:2d}. {name} ({cat}) - {district}")

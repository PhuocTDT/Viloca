import os, sys, shutil
from PIL import Image
import openpyxl
from openpyxl.drawing.image import Image as OpenpyxlImage

sys.stdout.reconfigure(encoding='utf-8')

TPHCM_DIR = 'd:/Desktop/Viloca/tphcm'
JPG_DIR = 'd:/Desktop/Viloca/jpg'
TPHCM_JPG_DIR = 'd:/Desktop/Viloca/tphcm_jpg'
THUMBS_DIR = 'd:/Desktop/Viloca/jpg_thumbs'
EXCEL_FILES = [
    'd:/Desktop/Viloca/Danh_Sach_45_Dia_Diem_Du_Lich_TPHCM.xlsx',
    'd:/Desktop/Viloca/Danh_Sach_45_Dia_Diem_Du_Lich_TP_Ho_Chi_Minh.xlsx'
]

os.makedirs(JPG_DIR, exist_ok=True)
os.makedirs(TPHCM_JPG_DIR, exist_ok=True)
os.makedirs(THUMBS_DIR, exist_ok=True)

MAPPINGS = [
    (1, '01_Dinh_Doc_Lap', 'Reunification_Palace_-_Saigon_(5539083573).jpg'),
    (2, '02_Nha_tho_Duc_Ba_Sai_Gon', 'Basílica_de_Nuestra_Señora,_Ciudad_Ho_Chi_Minh,_Vietnam,_2013-08-14,_DD_03.jfif'),
    (3, '03_Buu_dien_Trung_tam_Sai_Gon', 'Ho_Chi_Minh_City,_Central_Post_Office,_2020-01_CN-01.jpg'),
    (4, '04_Nha_hat_Thanh_pho', 'Ópera,_Ciudad_Ho_Chi_Minh,_Vietnam,_2013-08-14,_DD_02.jfif'),
    (5, '05_Tru_so_HDND_UBND_TPHCM', 'Ho_Chi_Minh_City,_City_Hall,_2020-01_CN-01.jpg'),
    (6, '06_Cho_Ben_Thanh', 'Ben_Thanh,_Ciudad_Ho_Chi_Minh,_Vietnam,_2013-08-14,_DD_01.jfif'),
    (7, '07_Bao_tang_Chung_tich_Chien_tranh', 'OIP.webp'),
    (8, '08_Bao_tang_My_thuat_TPHCM', 'Bao-tang-My-thuat-TP-HCM-ivivu.jpg'),
    (9, '09_Bao_tang_Lich_su_TPHCM', '10-dia-diem-du-lich-thanh-pho-ho-chi-minh-noi-tieng-nhat-5fa11376c8f9b.jpg'),
    (10, '10_Bao_tang_TP_Ho_Chi_Minh_Dinh_Gia_Long', '10_Bao_tang_TP_Ho_Chi_Minh_Dinh_Gia_Long.jpg'),
    (11, '11_Ben_Nha_Rong_Bao_tang_Ho_Chi_Minh', '11_Ben_Nha_Rong_Bao_tang_Ho_Chi_Minh.jpg'),
    (12, '12_Toa_nha_Landmark_81', 'dai-quan-sat-landmark-81-skyview-07-1716134940.jpg'),
    (13, '13_Cong_vien_Vinhomes_Central_Park', 'hai-thap-vong-canh-an-tuong-doc-bo-song-sai-gon-mang-den-tam-nhin-bao-quat-cong-vien-vinhomes-central-park.jpg'),
    (14, '14_Toa_thap_Bitexco_Financial_Tower', '155.jpg'),
    (15, '15_Pho_di_bo_Nguyen_Hue', 'pho-di-bo-nguyen-hue-o-dau-4.jpg'),
    (16, '16_Chung_cu_Ca_phe_42_Nguyen_Hue', '42-Nguyen-Hue-Finished-2755-FB-final.jpg'),
    (17, '17_Duong_Sach_TPHCM', 'duong-sach-thanh-pho-ho-chi-minh-1.jpg'),
    (18, '18_Pho_di_bo_Bui_Vien', 'vietgoing_mfj2112173210.webp'),
    (19, '19_Ben_Bach_Dang_Saigon_Waterbus', 'xe-buyt-song-sai-gon.jpg'),
    (20, '20_Cau_Ba_Son', 'cau-ba-son-thiet-ke-1736564498.jpg'),
    (21, '21_Cho_hoa_Ho_Thi_Ky', 'cho-hoa-Ho-Thi-Ky-5.jpg'),
    (22, '22_Cho_Tan_Dinh', 'edit_photo_tbsuvufcrk-1778660399.webp'),
    (23, '23_Pho_long_den_Luong_Nhu_Hoc', 'images.jfif'),
    (24, '24_Thao_Cam_Vien_Sai_Gon', 'images (1).jfif'),
    (25, '25_Cong_vien_Dam_Sen', 'HON-SEN-KHO-02-min.jpg'),
    (26, '26_Cau_Anh_Sao_Phu_My_Hung', 'cau-anh-sao-2.jpg'),
    (27, '27_Dia_dao_Cu_Chi', 'dia-dao-cu-chi-850.jpeg'),
    (28, '28_Rung_ngap_man_Can_Gio_Dao_Khi', '5352_Rung-ngap-man-Can-Gio-1-800x494.jpg'),
    (29, '29_Chua_Buu_Long', '25_Chua_Buu_Long.jpg'),
    (30, '30_Chua_Ngoc_Hoang', '23_Chua_Ngoc_Hoang_Jade_Emperor.jpg'),
    (31, '31_Chua_Vinh_Nghiem', '24_Chua_Vinh_Nghiem.jpg'),
    (32, '32_Chua_Ba_Thien_Hau_Cho_Lon', '26_Chua_Ba_Thien_Hau_Cho_Lon.jpg'),
    (33, '33_Hoi_quan_Nghia_An_Chua_Ong', '27_Hoi_quan_Nghia_An_Chua_Ong.jpg'),
    (34, '34_Cho_Binh_Tay_Cho_Lon', '34_Cho_Binh_Tay_Cho_Lon.jpg'),
    (35, '35_Cau_Mong_Saigon', '21_Cau_Mong_Saigon.jpg'),
    (36, '36_Chua_Giac_Lam', '28_Chua_Giac_Lam.jpg'),
    (37, '37_Lang_Ong_Ba_Chieu_Le_Van_Duyet', '29_Lang_Ong_Ba_Chieu_Le_Van_Duyet.jpg'),
    (38, '38_Chua_Phap_Hoa_Nhieu_Loc', '30_Chua_Phap_Hoa_Nhieu_Loc.jpg'),
    (39, '39_Den_Mariamman_Chua_Ba_An', '31_Den_Mariamman_Chua_Ba_An.jpg'),
    (40, '40_Thanh_duong_Jamia_Al_Musulman', '32_Thanh_duong_Jamia_Al_Musulman.jpg'),
    (41, '41_Nha_tho_Cha_Tam_Cho_Lon', '33_Nha_tho_Cha_Tam_Cho_Lon.jpg'),
    (42, '42_Pho_do_co_Le_Cong_Kieu', '38_Pho_do_co_Le_Cong_Kieu.jpg'),
    (43, '43_KDL_Suoi_Tien', '41_KDL_Suoi_Tien.jpg'),
    (44, '44_Nong_trang_Xanh_Green_Noen_Cu_Chi', '45_Nong_trang_Xanh_Green_Noen_Cu_Chi.jpg'),
    (45, '45_Pho_am_thuc_Ha_Ton_Quyen', '45_Pho_am_thuc_Ha_Ton_Quyen.jpg')
]

def convert_to_jpg(src_path, dst_path):
    try:
        with Image.open(src_path) as img:
            if img.mode in ('RGBA', 'LA', 'P'):
                rgb_img = Image.new('RGB', img.size, (255, 255, 255))
                if img.mode == 'P':
                    img = img.convert('RGBA')
                rgb_img.paste(img, mask=img.split()[3] if len(img.split()) > 3 else None)
                rgb_img.save(dst_path, 'JPEG', quality=95)
            else:
                img.convert('RGB').save(dst_path, 'JPEG', quality=95)
        return True
    except Exception as e:
        print(f"Error converting {src_path}: {e}")
        return False

def make_thumbnail(src_path, dst_thumb_path, target_size=(150, 90)):
    try:
        with Image.open(src_path) as img:
            img = img.convert('RGB')
            src_w, src_h = img.size
            tgt_w, tgt_h = target_size
            src_aspect = src_w / src_h
            tgt_aspect = tgt_w / tgt_h

            if src_aspect > tgt_aspect:
                new_h = tgt_h
                new_w = int(src_w * (tgt_h / src_h))
                img_resized = img.resize((new_w, new_h), Image.Resampling.LANCZOS)
                left = (new_w - tgt_w) // 2
                img_cropped = img_resized.crop((left, 0, left + tgt_w, tgt_h))
            else:
                new_w = tgt_w
                new_h = int(src_h * (tgt_w / src_w))
                img_resized = img.resize((new_w, new_h), Image.Resampling.LANCZOS)
                top = (new_h - tgt_h) // 2
                img_cropped = img_resized.crop((0, top, tgt_w, top + tgt_h))

            img_cropped.save(dst_thumb_path, 'JPEG', quality=92)
        return True
    except Exception as e:
        print(f"Error making thumb for {src_path}: {e}")
        return False

# Clear jpg/ and tphcm_jpg/ directories first
for f in os.listdir(JPG_DIR):
    os.remove(os.path.join(JPG_DIR, f))
for f in os.listdir(TPHCM_JPG_DIR):
    os.remove(os.path.join(TPHCM_JPG_DIR, f))

print("=== BẮT ĐẦU CHUYỂN ĐỔI & ĐỔI TÊN ẢNH TP.HCM THEO THỨ TỰ ===")

primary_thumbs = {}

for stt, clean_name, src_file in MAPPINGS:
    src_p = os.path.join(TPHCM_DIR, src_file)
    dst_jpg = os.path.join(JPG_DIR, f"{clean_name}.jpg")
    dst_tphcm_jpg = os.path.join(TPHCM_JPG_DIR, f"{clean_name}.jpg")
    thumb_dst = os.path.join(THUMBS_DIR, f"{clean_name}_thumb.jpg")

    if os.path.exists(src_p):
        convert_to_jpg(src_p, dst_jpg)
        shutil.copy(dst_jpg, dst_tphcm_jpg)
        make_thumbnail(dst_jpg, thumb_dst)
        primary_thumbs[stt] = thumb_dst
        print(f"[{stt:02d}] {src_file} -> {clean_name}.jpg")
    else:
        print(f"[{stt:02d}] CẢNH BÁO: Không tìm thấy {src_p}")

print(f"\n=> Đã chuyển đổi hoàn tất {len(primary_thumbs)}/45 ảnh vào thư mục jpg/ và tphcm_jpg/!")

# Update Excel files
for ex_path in EXCEL_FILES:
    if os.path.exists(ex_path):
        print(f"\n=== CẬP NHẬT FILE EXCEL: {os.path.basename(ex_path)} ===")
        wb = openpyxl.load_workbook(ex_path)
        ws = wb.active

        # Clear existing images
        ws._images.clear()

        # Map row index
        row_for_stt = {}
        for r in range(4, ws.max_row + 1):
            val = ws.cell(r, 1).value
            if val is not None:
                try:
                    row_for_stt[int(val)] = r
                except:
                    pass

        inserted = 0
        for stt, thumb_p in primary_thumbs.items():
            if stt in row_for_stt and os.path.exists(thumb_p):
                row_idx = row_for_stt[stt]
                img = OpenpyxlImage(thumb_p)
                img.width = 150
                img.height = 90
                cell_ref = f"J{row_idx}"
                ws.add_image(img, cell_ref)
                ws.row_dimensions[row_idx].height = 82
                inserted += 1

        ws.column_dimensions['J'].width = 24
        wb.save(ex_path)
        print(f"=> Đã nhúng thành công {inserted} ảnh vào {os.path.basename(ex_path)}")

# Also sync to TP_Ho_Chi_Minh.xlsx if needed
shutil.copy('d:/Desktop/Viloca/Danh_Sach_45_Dia_Diem_Du_Lich_TPHCM.xlsx', 'd:/Desktop/Viloca/TP_Ho_Chi_Minh.xlsx')
print("=> Đã đồng bộ sang TP_Ho_Chi_Minh.xlsx")
